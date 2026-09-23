"""
Извлечение текста из файлов — по той же логике, что default-engine
Open WebUI (backend/open_webui/retrieval/loaders/main.py):
PyPDFLoader, Docx2txtLoader, TextLoader, CSV, Excel (pandas fallback), PPTX.
"""

from __future__ import annotations

import logging
import zipfile
from pathlib import Path
from typing import Any

import ftfy

from ...config import ALLOWED_EXTENSIONS

log = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".xml",
    ".csv",
    ".log",
    ".ini",
    ".yml",
    ".yaml",
    ".toml",
    ".py",
    ".js",
    ".ts",
    ".css",
    ".html",
    ".htm",
    ".rtf",
}


def _fix(text: str) -> str:
    return ftfy.fix_text(text or "").strip()


def _detect_encoding(path: Path) -> str:
    raw = path.read_bytes()[:65536]
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _extract_legacy_doc(path: Path) -> str:
    """Extract text from legacy .doc via antiword/catdoc (installed in image)."""
    import shutil
    import subprocess

    errors: list[str] = []
    for cmd in (
        ["antiword", "-m", "UTF-8.txt", str(path)],
        ["antiword", str(path)],
        ["catdoc", "-d", "utf-8", str(path)],
        ["catdoc", str(path)],
    ):
        binary = cmd[0]
        if not shutil.which(binary):
            errors.append(f"{binary} not installed")
            continue
        try:
            proc = subprocess.run(
                cmd,
                check=False,
                capture_output=True,
                timeout=60,
            )
            if proc.returncode == 0 and (proc.stdout or b"").strip():
                text = proc.stdout.decode("utf-8", errors="replace")
                if not text.strip():
                    text = proc.stdout.decode("cp1251", errors="replace")
                return text
            err = (proc.stderr or b"").decode("utf-8", errors="replace")[:200]
            errors.append(f"{' '.join(cmd)} -> rc={proc.returncode} {err}")
        except Exception as e:  # noqa: BLE001
            errors.append(f"{binary}: {e}")

    raise ValueError(
        "Cannot read legacy .doc ("
        + "; ".join(errors[:4])
        + "). Convert to .docx or install antiword/catdoc."
    )


def extract_file(path: Path) -> dict[str, Any]:
    """Вернуть извлечённый текст одного файла (как после upload в чат OWUI)."""
    if not path.is_file():
        raise FileNotFoundError(f"Not a file: {path}")

    ext = path.suffix.lower()
    if ext not in ALLOWED_EXTENSIONS and ext not in TEXT_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext == ".zip":
        return _extract_zip(path)

    text = _extract_by_ext(path, ext)
    text = _fix(text)
    return {
        "path": str(path),
        "filename": path.name,
        "extension": ext,
        "chars": len(text),
        "content": text,
    }


def _extract_by_ext(path: Path, ext: str) -> str:
    if ext == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader

        docs = PyPDFLoader(str(path), mode="page").load()
        return "\n\n".join(d.page_content for d in docs)

    if ext == ".docx":
        from langchain_community.document_loaders import Docx2txtLoader

        docs = Docx2txtLoader(str(path)).load()
        return "\n\n".join(d.page_content for d in docs)

    if ext in {".html", ".htm"}:
        from langchain_community.document_loaders import BSHTMLLoader

        try:
            docs = BSHTMLLoader(str(path), open_encoding="utf-8").load()
        except Exception:
            docs = BSHTMLLoader(str(path), open_encoding="unicode_escape").load()
        return "\n\n".join(d.page_content for d in docs)

    if ext == ".csv":
        enc = _detect_encoding(path)
        import pandas as pd

        # Как при «простом» чтении таблицы: сохранить строки целиком
        for sep in (";", ",", "\t"):
            try:
                df = pd.read_csv(path, encoding=enc, sep=sep)
                if df.shape[1] > 1:
                    break
            except Exception:
                continue
        else:
            df = pd.read_csv(path, encoding=enc, sep=None, engine="python")
        summary = (
            f"Table: {len(df) + 1} rows incl. header; {len(df)} data rows; "
            f"{len(df.columns)} columns: {', '.join(map(str, df.columns))}."
        )
        return f"{summary}\n\n{df.to_string(index=False)}"

    if ext in {".xlsx", ".xls"}:
        import pandas as pd

        parts: list[str] = []
        engine = "xlrd" if ext == ".xls" else "openpyxl"
        try:
            xls = pd.ExcelFile(path, engine=engine)
        except Exception:
            # fallback: let pandas choose
            xls = pd.ExcelFile(path)
        for sheet in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet)
            parts.append(f"Sheet: {sheet}\n{df.to_string(index=False)}")
        return "\n\n".join(parts)

    if ext in {".pptx", ".ppt"}:
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [
                    shape.text_frame.text
                    for shape in slide.shapes
                    if shape.has_text_frame and shape.text_frame.text.strip()
                ]
                if texts:
                    parts.append(f"Slide {i}:\n" + "\n".join(texts))
            return "\n\n".join(parts)
        except Exception as e:
            raise ValueError(f"Cannot read PowerPoint: {e}") from e

    if ext == ".doc":
        return _extract_legacy_doc(path)

    # text-like
    enc = _detect_encoding(path)
    return path.read_text(encoding=enc, errors="replace")


def _extract_zip(path: Path) -> dict[str, Any]:
    """Распаковать zip во временный смысл: извлечь текст из вложенных файлов."""
    parts: list[str] = []
    files_meta: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = info.filename
            # skip macOS junk / hidden
            if "__MACOSX" in name or Path(name).name.startswith("."):
                continue
            ext = Path(name).suffix.lower()
            if ext not in ALLOWED_EXTENSIONS and ext not in TEXT_EXTENSIONS:
                files_meta.append(
                    {"filename": name, "skipped": True, "reason": f"unsupported {ext}"}
                )
                continue
            if ext == ".zip":
                files_meta.append(
                    {"filename": name, "skipped": True, "reason": "nested zip"}
                )
                continue
            try:
                raw = zf.read(info)
                tmp = path.parent / f".__zip_extract_{Path(name).name}"
                try:
                    tmp.write_bytes(raw)
                    extracted = _extract_by_ext(tmp, ext)
                    extracted = _fix(extracted)
                    parts.append(f"===== {name} =====\n{extracted}")
                    files_meta.append(
                        {"filename": name, "chars": len(extracted), "skipped": False}
                    )
                finally:
                    if tmp.exists():
                        tmp.unlink()
            except Exception as e:
                files_meta.append(
                    {"filename": name, "skipped": True, "reason": str(e)}
                )

    content = _fix("\n\n".join(parts))
    return {
        "path": str(path),
        "filename": path.name,
        "extension": ".zip",
        "chars": len(content),
        "content": content,
        "zip_files": files_meta,
    }
