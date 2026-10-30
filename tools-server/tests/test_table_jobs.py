import hashlib
import threading

import pytest
from openpyxl import Workbook, load_workbook

from app import table_tool as tables
from app.table_tool import jobs
from app import document_tool, excel_tool


@pytest.fixture
def setup_job(tmp_path, monkeypatch):
    for module in (tables, jobs, document_tool, excel_tool):
        monkeypatch.setattr(module, 'DATA_ROOT', tmp_path)
    monkeypatch.setattr(jobs, '_jobs', {})
    monkeypatch.setattr(jobs, '_llm_config', lambda: ('url', 'key', 'model'))
    # Deterministic execution of the actual worker, without remote model calls.
    class InlineThread:
        def __init__(self, target, args, **kwargs):
            self.target, self.args = target, args
        def start(self):
            self.target(*self.args)
    monkeypatch.setattr(jobs.threading, 'Thread', InlineThread)
    # ThreadPoolExecutor needs real threads; execute batches sequentially in tests.
    class Pool:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def map(self, fn, rows): return map(fn, rows)
    monkeypatch.setattr(jobs, 'ThreadPoolExecutor', Pool)
    src = tmp_path / 'exports' / 'input.xlsx'
    src.parent.mkdir()
    wb = Workbook()
    wb.active.append(['Article', 'Description', 'Type'])
    wb.active.append(['123', None, None])
    wb.active.append(['456', None, None])
    wb.save(src)
    return tmp_path, src


def start():
    return jobs.start_job('exports/input.xlsx', 'Fill', ['A'],
                          {'B': {}, 'C': {'allowed': ['valid']}}, header_row=1, batch_size=1)


def test_failed_calls_are_failed(setup_job, monkeypatch):
    def fail(*args): raise RuntimeError('offline')
    monkeypatch.setattr(jobs, '_run_batch', fail)
    status = start()
    assert status['state'] == 'failed'
    assert status['rows_filled'] == status['cells_written'] == 0
    assert status['errors_total'] == 2


def test_rejected_cells_count_and_independent_outputs(setup_job, monkeypatch):
    root, src = setup_job
    before = hashlib.sha256(src.read_bytes()).digest()
    monkeypatch.setattr(jobs, '_run_batch', lambda j, batch, *args: (
        [{'row': batch[0]['row'], 'values': {'B': 'description', 'C': 'invalid'}, 'confidence': 'high'}], {}))
    one, two = start(), start()
    assert one['state'] == 'partial'
    assert one['rows_filled'] == 0
    assert one['rows_written'] == 2 and one['cells_written'] == 2
    assert one['output'] != two['output'] != 'exports/input.xlsx'
    assert hashlib.sha256(src.read_bytes()).digest() == before
    wb = load_workbook(root / one['output'])
    assert wb.active['B2'].value == 'description' and wb.active['C2'].value is None
    wb.close()


def test_complete_and_interrupted_checkpoint(setup_job, monkeypatch):
    monkeypatch.setattr(jobs, '_run_batch', lambda j, batch, *args: (
        [{'row': batch[0]['row'], 'values': {'B': 'description', 'C': 'valid'}, 'confidence': 'high'}], {}))
    status = start()
    assert status['state'] == 'done' and status['rows_filled'] == 2
    job = jobs._jobs.pop(status['job_id'])
    job['state'] = 'running'
    jobs._save_status(job)
    recovered = jobs.job_status(status['job_id'])
    assert recovered['state'] == 'interrupted'
    assert recovered['output'] == status['output'] and recovered['rows_filled'] == 2
    assert jobs.job_status(status['job_id'])['errors_total'] == 1


def test_atomic_save_preserves_previous_file_on_failure(setup_job, monkeypatch):
    root, src = setup_job
    before = src.read_bytes()
    def fail(self, path):
        from pathlib import Path
        Path(path).write_bytes(b'incomplete zip')
        raise OSError('disk error')
    monkeypatch.setattr('openpyxl.workbook.workbook.Workbook.save', fail)
    with pytest.raises(OSError):
        tables.fill_table('exports/input.xlsx', [{'row': 2, 'values': {'B': 'new'}}], header_row=1)
    assert src.read_bytes() == before
    assert not list(src.parent.glob('.table-*'))
