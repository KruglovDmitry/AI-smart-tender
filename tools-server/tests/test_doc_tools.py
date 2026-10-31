"""Инструменты больших документов и таблиц: поиск артикулов, чтение страниц, запрос и заполнение Excel."""

from __future__ import annotations

import importlib

import pytest
from openpyxl import Workbook, load_workbook


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    from app import config

    importlib.reload(config)
    import app.document_tool as dt
    import app.excel_tool as et
    import app.doc_index as di
    import app.table_tool as tt
    import app.main as main

    for m in (dt, et, di, tt, main):
        importlib.reload(m)
    from fastapi.testclient import TestClient

    return TestClient(main.app), tmp_path


def _catalog(path):
    import pymupdf

    doc = pymupdf.open()
    for text in (
        "СЕРИЯ 55\nТип 55.32 - 2 CO 10 A\nКодировка: 55.32.9.024.0000",
        "Тип 55.32 катушки AC и DC\nПример: 55.34.9.220.9202",
        "Твердотельное реле 34.81.7.012.9024 питание 12В DC",
    ):
        doc.new_page().insert_text((50, 72), text, fontname="helv")
    doc.save(path)


def test_search_code_exact_and_prefix(client):
    c, root = client
    (root / "catalogs").mkdir()
    _catalog(root / "catalogs" / "cat.pdf")

    r = c.post("/search_document", json={"path": "catalogs/cat.pdf", "query": "348170129024"}).json()
    assert r["match"] == "exact" and [x["page"] for x in r["results"]] == [3]

    # Полного кода нет; самый длинный напечатанный префикс 55.32.9 есть только в кодировке на стр. 1.
    r = c.post("/search_document", json={"path": "catalogs/cat.pdf", "query": "55.32.9.012.0000"}).json()
    assert r["match"] == "prefix" and r["matched"] == "55329"
    assert [x["page"] for x in r["results"]] == [1]

    r = c.post("/search_document", json={"path": "catalogs/cat.pdf", "query": "55.32.7.110.0000"}).json()
    assert r["matched"] == "5532" and {x["page"] for x in r["results"]} == {1, 2}


def test_read_pages_and_info(client):
    c, root = client
    (root / "catalogs").mkdir()
    _catalog(root / "catalogs" / "cat.pdf")
    info = c.post("/document_info", json={"path": "catalogs/cat.pdf"}).json()
    assert info["pages"] == 3 and info["broken_pages"] == []
    r = c.post("/read_pages", json={"path": "catalogs/cat.pdf", "pages": "2-3", "include_tables": False}).json()
    assert [p["page"] for p in r["pages"]] == [2, 3] and "34.81.7.012.9024" in r["pages"][1]["text"]
    assert c.post("/read_pages", json={"path": "catalogs/cat.pdf", "pages": "9"}).status_code == 400


def test_owui_loader_rtf_text_not_markup(client):
    c, root = client
    word = "Техническое задание".encode("cp1251")
    rtf = r"{\rtf1\ansi\ansicpg1251{\fonttbl{\f0 Times New Roman;}}\f0 " + "".join(f"\\'{b:02x}" for b in word) + r"\par}"
    r = c.put("/process", content=rtf.encode("ascii"), headers={"X-Filename": "tz.rtf"})
    text = r.json()[0]["page_content"]
    assert "Техническое задание" in text and "\\rtf1" not in text


def test_owui_loader_extensionless_pdf(client):
    """Имя без расширения (как с ЕИС) — тип по содержимому."""
    import pymupdf

    c, _root = client
    buf = __import__("io").BytesIO()
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 72), "Delivery term 30 days", fontname="helv")
    doc.save(buf)
    doc.close()
    r = c.put(
        "/process",
        content=buf.getvalue(),
        headers={"X-Filename": "I. Obschie svedeniya (tovar)"},
    )
    assert r.status_code == 200
    text = r.json()[0]["page_content"]
    assert "Delivery" in text or "30" in text


def test_extract_file_sniffs_extensionless_pdf(tmp_path, monkeypatch):
    import importlib

    import pymupdf

    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    from app import config

    importlib.reload(config)
    import app.document_tool.extract as ex

    importlib.reload(ex)
    path = tmp_path / "I. Obschie svedeniya"
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 72), "hello pdf", fontname="helv")
    doc.save(path)
    doc.close()
    out = ex.extract_file(path)
    assert out["extension"] == ".pdf" and "hello" in out["content"]


def _template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "IMPORT"
    ws.append(["Код ТНВЭД", "Полное наименование", "Вид товара", "Артикул"])
    ws.append(["Tnved", "2478", "12", "13914"])
    ws.append(["8536490000", None, None, 553290240000])
    ws.append(["8536490000", None, None, 405180600000])
    ref = wb.create_sheet("Справочники")
    ref.append(["Вид товара", "РЕЛЕ ПРОМЕЖУТОЧНОЕ МИНИАТЮРНОЕ"])
    wb.save(path)


def test_query_and_fill_table(client):
    c, root = client
    (root / "catalogs").mkdir()
    _template(root / "catalogs" / "t.xlsx")

    d = c.post("/describe_table", json={"path": "catalogs/t.xlsx", "sheet": "IMPORT"}).json()
    assert d["details"][0]["header_row"] == 1

    q = c.post(
        "/query_table",
        json={"path": "catalogs/t.xlsx", "data_start_row": 3, "columns": ["Артикул"], "filters": [{"column": "B", "op": "empty"}]},
    ).json()
    assert [r["row"] for r in q["rows"]] == [3, 4]

    f = c.post(
        "/fill_table",
        json={
            "path": "catalogs/t.xlsx",
            "key_column": "D",
            "updates": [
                {"key": "553290240000", "values": {"B": "Реле 2CO", "Вид товара": "РЕЛЕ ПРОМЕЖУТОЧНОЕ МИНИАТЮРНОЕ"}},
                {"key": "405180600000", "values": {"C": "ВЫДУМАННЫЙ ВИД"}},
            ],
            "allowed_values": {"C": ["РЕЛЕ ПРОМЕЖУТОЧНОЕ МИНИАТЮРНОЕ"]},
        },
    ).json()
    assert f["cells_written"] == 2 and f["problems_total"] == 1 and f["path"].startswith("exports/")

    wb = load_workbook(root / f["path"])
    assert wb["IMPORT"]["B3"].value == "Реле 2CO" and wb["IMPORT"]["C4"].value is None
    assert wb.sheetnames == ["IMPORT", "Справочники"]
    # Исходный шаблон не тронут
    assert load_workbook(root / "catalogs" / "t.xlsx")["IMPORT"]["B3"].value is None
