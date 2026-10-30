"""Работа с большими таблицами без чтения их в контекст модели.

Модель получает структуру (листы, колонки, примеры), запрашивает нужные строки фильтром
и записывает значения в копию исходного файла — оформление, листы и справочники шаблона сохраняются.
"""

from __future__ import annotations

import re
import shutil
import os
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from ..config import DATA_ROOT
from ..document_tool import resolve_under_data
from ..excel_tool import export_dest, make_download_url

MAX_QUERY_ROWS = 200
MAX_FILL_UPDATES = 2000
MARK_FILL = PatternFill("solid", fgColor="FFF2CC")  # бледно-жёлтый: «проверить руками»
HEADER_SCAN_ROWS = 15


def _val(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


class Sheet:
    """Лист как сетка значений (строки с 1), загруженная один раз."""

    def __init__(self, name: str, rows: list[list[Any]]):
        self.name = name
        self.rows = rows
        self.width = max((len(r) for r in rows), default=0)

    def cell(self, r: int, c: int) -> Any:
        row = self.rows[r - 1] if 0 < r <= len(self.rows) else []
        return row[c - 1] if 0 < c <= len(row) else None

    def detect_header(self) -> int:
        """Первая строка, где заполнено больше половины колонок, и в основном текстом."""
        best, best_score = 1, -1.0
        for r in range(1, min(len(self.rows), HEADER_SCAN_ROWS) + 1):
            vals = [v for v in self.rows[r - 1] if not _blank(v)]
            if not vals:
                continue
            texts = sum(isinstance(v, str) and not re.fullmatch(r"[\d\s.,-]+", v) for v in vals)
            score = len(vals) / max(1, self.width) + texts / len(vals)
            if score > best_score + 0.25:
                best, best_score = r, score
        return best


def _load(path: Path) -> dict[str, Sheet]:
    ext = path.suffix.lower()
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return {ws.title: Sheet(ws.title, [list(map(_val, r)) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets}
        finally:
            wb.close()
    import pandas as pd

    if ext == ".csv":
        from ..document_tool.extract import _detect_encoding

        df = pd.read_csv(path, header=None, sep=None, engine="python", encoding=_detect_encoding(path), dtype=object)
        frames = {"csv": df}
    elif ext == ".xls":
        frames = pd.read_excel(path, sheet_name=None, header=None, dtype=object, engine="xlrd")
    else:
        raise ValueError(f"Not a table file: {ext} (use .xlsx, .xls, .csv)")
    out = {}
    for name, df in frames.items():
        rows = [[None if (isinstance(v, float) and v != v) else _val(v) for v in r] for r in df.itertuples(index=False)]
        out[str(name)] = Sheet(str(name), rows)
    return out


def _pick_sheet(sheets: dict[str, Sheet], sheet: str | None) -> Sheet:
    if not sheet:
        return next(iter(sheets.values()))
    if sheet in sheets:
        return sheets[sheet]
    low = {k.lower(): v for k, v in sheets.items()}
    if sheet.lower() in low:
        return low[sheet.lower()]
    raise ValueError(f"Sheet {sheet!r} not found; sheets: {list(sheets)}")


def _headers(sh: Sheet, header_row: int) -> list[str]:
    return [("" if _blank(sh.cell(header_row, c)) else str(sh.cell(header_row, c)).strip()) for c in range(1, sh.width + 1)]


def _resolve_column(sh: Sheet, header_row: int, ref: str) -> int:
    """Колонка по букве (B, AA), номеру или тексту заголовка (точно, затем по вхождению)."""
    ref = str(ref).strip()
    if re.fullmatch(r"[A-Za-z]{1,3}", ref):
        return column_index_from_string(ref.upper())
    if ref.isdigit():
        return int(ref)
    heads = _headers(sh, header_row)
    for i, h in enumerate(heads, 1):
        if h.lower() == ref.lower():
            return i
    found = [i for i, h in enumerate(heads, 1) if ref.lower() in h.lower()]
    if len(found) == 1:
        return found[0]
    if found:
        raise ValueError(f"Column {ref!r} is ambiguous: {[get_column_letter(i) + ' ' + heads[i - 1] for i in found]}")
    raise ValueError(f"Column {ref!r} not found in header row {header_row}")


def _col_label(sh: Sheet, header_row: int, c: int) -> str:
    h = _headers(sh, header_row)[c - 1] if c <= sh.width else ""
    return f"{get_column_letter(c)}" + (f" ({h})" if h else "")


def describe_table(path: str, sheet: str | None = None, header_row: int | None = None, sample_rows: int = 3) -> dict[str, Any]:
    file_path = resolve_under_data(path)
    sheets = _load(file_path)
    rel = str(file_path.relative_to(DATA_ROOT)).replace("\\", "/")
    targets = [_pick_sheet(sheets, sheet)] if sheet else list(sheets.values())
    out = []
    for sh in targets:
        hr = header_row or sh.detect_header()
        heads = _headers(sh, hr)
        data = sh.rows[hr:]
        last = max((i for i, r in enumerate(data, 1) if any(not _blank(v) for v in r)), default=0)
        data = data[:last]
        cols = []
        for c in range(1, sh.width + 1):
            vals = [r[c - 1] for r in data if c <= len(r) and not _blank(r[c - 1])]
            if not vals and not heads[c - 1]:
                continue
            distinct = list(dict.fromkeys(map(str, vals)))
            col = {"column": get_column_letter(c), "header": heads[c - 1], "filled": len(vals)}
            if 0 < len(distinct) <= 12:
                col["values"] = distinct
            else:
                col["distinct"] = len(distinct)
                col["examples"] = distinct[:3]
            cols.append(col)
        out.append(
            {
                "sheet": sh.name,
                "header_row": hr,
                "rows_above_header": [
                    [_val(v) for v in r[:30]] for r in sh.rows[: hr - 1]
                ][:6],
                "rows_below_header_first": [
                    {"row": hr + i, "values": [_val(v) for v in r[:30]]} for i, r in enumerate(data[: min(sample_rows, 10)], 1)
                ],
                "data_rows": len(data),
                "columns": cols,
            }
        )
    return {
        "path": rel,
        "sheets": [s.name for s in sheets.values()],
        "details": out,
        "note": (
            "Do not try to read big tables whole. Use query_table with filters/offset to get rows "
            "(max 200 per call) and fill_table to write values back into a copy of the file. "
            "If the real header is not the detected one, pass header_row explicitly; if service rows follow "
            "the header (ids, types, hints), pass data_start_row to query_table."
        ),
    }


_OPS = {"eq", "ne", "contains", "not_contains", "empty", "not_empty", "in", "startswith", "gt", "lt"}


def _match(v: Any, op: str, target: Any) -> bool:
    s = "" if v is None else str(v).strip().lower()
    t = "" if target is None else str(target).strip().lower()
    if op == "empty":
        return _blank(v)
    if op == "not_empty":
        return not _blank(v)
    if op == "eq":
        return s == t
    if op == "ne":
        return s != t
    if op == "contains":
        return t in s
    if op == "not_contains":
        return t not in s
    if op == "startswith":
        return s.startswith(t)
    if op == "in":
        items = target if isinstance(target, list) else str(target).split("|")
        return s in {str(x).strip().lower() for x in items}
    if op in {"gt", "lt"}:
        try:
            a, b = float(str(v).replace(",", ".")), float(str(target).replace(",", "."))
        except (TypeError, ValueError):
            return False
        return a > b if op == "gt" else a < b
    raise ValueError(f"Unknown op {op!r}; use one of {sorted(_OPS)}")


def query_table(
    path: str,
    sheet: str | None = None,
    header_row: int | None = None,
    data_start_row: int | None = None,
    columns: list[str] | None = None,
    filters: list[dict[str, Any]] | None = None,
    offset: int = 0,
    limit: int = 50,
    distinct: str | None = None,
) -> dict[str, Any]:
    file_path = resolve_under_data(path)
    sh = _pick_sheet(_load(file_path), sheet)
    hr = header_row or sh.detect_header()
    conds = []
    for f in filters or []:
        op = str(f.get("op", "eq"))
        if op not in _OPS:
            raise ValueError(f"Unknown op {op!r}; use one of {sorted(_OPS)}")
        conds.append((_resolve_column(sh, hr, f["column"]), op, f.get("value")))
    matched = [
        r for r in range(max(hr + 1, data_start_row or 0), len(sh.rows) + 1)
        if any(not _blank(v) for v in sh.rows[r - 1]) and all(_match(sh.cell(r, c), op, t) for c, op, t in conds)
    ]
    if distinct:
        c = _resolve_column(sh, hr, distinct)
        counts: dict[str, int] = {}
        for r in matched:
            key = "" if _blank(sh.cell(r, c)) else str(sh.cell(r, c))
            counts[key] = counts.get(key, 0) + 1
        top = sorted(counts.items(), key=lambda kv: -kv[1])
        return {
            "sheet": sh.name, "column": _col_label(sh, hr, c), "matched_rows": len(matched),
            "distinct": len(top), "values": [{"value": k, "count": n} for k, n in top[:MAX_QUERY_ROWS]],
        }
    cols = [_resolve_column(sh, hr, c) for c in columns] if columns else [
        c for c in range(1, sh.width + 1) if _headers(sh, hr)[c - 1] or any(not _blank(sh.cell(r, c)) for r in matched[:50])
    ]
    limit = max(1, min(limit, MAX_QUERY_ROWS))
    page = matched[offset : offset + limit]
    return {
        "sheet": sh.name,
        "header_row": hr,
        "columns": [_col_label(sh, hr, c) for c in cols],
        "matched_rows": len(matched),
        "offset": offset,
        "returned": len(page),
        "rows": [{"row": r, "values": [sh.cell(r, c) for c in cols]} for r in page],
        **({"next_offset": offset + len(page)} if offset + len(page) < len(matched) else {}),
    }


def fill_table(
    path: str,
    updates: list[dict[str, Any]],
    sheet: str | None = None,
    header_row: int | None = None,
    key_column: str | None = None,
    allowed_values: dict[str, list[str]] | None = None,
    output: str | None = None,
    public_base_url: str = "",
) -> dict[str, Any]:
    """Записать значения в копию .xlsx (или продолжить уже созданную копию в exports/).

    updates: [{"row": 5, "values": {"D": "...", "Вид товара": "..."}}] или {"key": "553290240000", ...} при key_column.
    "mark": "<причина>" в update подсвечивает записанные ячейки и кладёт причину примечанием — что проверить руками.
    """
    if not updates:
        raise ValueError("No updates")
    if len(updates) > MAX_FILL_UPDATES:
        raise ValueError(f"Too many updates in one call (max {MAX_FILL_UPDATES})")
    src = resolve_under_data(path)
    if src.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("fill_table writes .xlsx/.xlsm only; convert .xls/.csv first")
    exports = (DATA_ROOT / "exports").resolve()
    if output:
        dest = resolve_under_data(output)
        dest.relative_to(exports)  # писать можно только в exports/
        if not dest.exists():
            shutil.copy2(src, dest)
    elif src.resolve().is_relative_to(exports):
        dest = src  # продолжаем заполнять уже созданную копию
    else:
        dest = export_dest(src.stem + "_filled.xlsx")
        shutil.copy2(src, dest)

    grid = _pick_sheet(_load(dest), sheet)
    hr = header_row or grid.detect_header()
    keys: dict[str, list[int]] = {}
    if key_column:
        kc = _resolve_column(grid, hr, key_column)
        for r in range(hr + 1, len(grid.rows) + 1):
            v = grid.cell(r, kc)
            if not _blank(v):
                keys.setdefault(re.sub(r"\s+", "", str(v)).lower(), []).append(r)
    allowed = {}
    for col, vals in (allowed_values or {}).items():
        allowed[_resolve_column(grid, hr, col)] = {str(v).strip().lower() for v in vals}

    wb = load_workbook(dest)
    ws = wb[grid.name]
    written, problems = 0, []
    written_rows: dict[int, set[str]] = {}
    col_cache: dict[str, int] = {}
    for i, u in enumerate(updates):
        if "row" in u and u["row"] is not None:
            rows = [int(u["row"])]
        elif key_column and u.get("key") is not None:
            rows = keys.get(re.sub(r"\s+", "", str(u["key"])).lower(), [])
            if not rows:
                problems.append({"update": i, "key": u["key"], "error": "key not found"})
                continue
        else:
            problems.append({"update": i, "error": "need row or key"})
            continue
        mark = str(u.get("mark") or "").strip()
        first_cell = True
        for ref, value in (u.get("values") or {}).items():
            try:
                c = col_cache.get(ref) or col_cache.setdefault(ref, _resolve_column(grid, hr, ref))
            except ValueError as e:
                problems.append({"update": i, "column": ref, "error": str(e)})
                continue
            if c in allowed and not _blank(value) and str(value).strip().lower() not in allowed[c]:
                problems.append({"update": i, "column": _col_label(grid, hr, c), "value": value, "error": "not in allowed values"})
                continue
            for r in rows:
                if r <= hr:
                    problems.append({"update": i, "row": r, "error": "row is in the header area"})
                    continue
                cell = ws.cell(row=r, column=c, value=value)
                written += 1
                written_rows.setdefault(r, set()).add(get_column_letter(c))
                if mark:
                    cell.fill = MARK_FILL
                    if first_cell:
                        cell.comment = Comment(mark[:500], "AI Smart Tender")
            first_cell = False
    # Replace only after the complete workbook is safely serialized.
    fd, temporary = tempfile.mkstemp(prefix=".table-", suffix=dest.suffix, dir=dest.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        wb.close()
        os.replace(temporary, dest)
    finally:
        wb.close()
        Path(temporary).unlink(missing_ok=True)
    rel = str(dest.relative_to(DATA_ROOT)).replace("\\", "/")
    return {
        "ok": not problems,
        "path": rel,
        "cells_written": written,
        "written_rows": {str(r): sorted(cols) for r, cols in written_rows.items()},
        "problems": problems[:50],
        "problems_total": len(problems),
        "download_url": make_download_url(rel, public_base_url),
        "note": "Pass this path back as `path` to continue filling the same copy. Give the user download_url as a markdown link.",
    }
